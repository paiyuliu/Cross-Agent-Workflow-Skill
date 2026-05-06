#!/usr/bin/env python3
"""Extract fuller text/table content from supported document formats."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import doc_inspect

try:
    from log_action import append_action
except Exception:  # pragma: no cover
    append_action = None


def extract_docx(path: Path) -> dict[str, Any]:
    try:
        import docx
    except ImportError:
        return doc_inspect.missing("docx", "python-docx")
    document = docx.Document(str(path))
    return {
        "ok": True,
        "type": "docx",
        "paragraphs": [p.text for p in document.paragraphs],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
    }


def extract_pptx(path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError:
        return doc_inspect.missing("pptx", "python-pptx")
    prs = Presentation(str(path))
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        blocks = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                blocks.append(shape.text)
        slides.append({"slide": idx, "text": blocks})
    return {"ok": True, "type": "pptx", "slides": slides}


def extract_xlsx(path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError:
        return doc_inspect.missing("openpyxl", "openpyxl")
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        sheets.append({"name": ws.title, "rows": rows})
    wb.close()
    return {"ok": True, "type": "xlsx", "sheets": sheets}


def extract_pdf(path: Path) -> dict[str, Any]:
    inspected = doc_inspect.inspect_pdf(path, max_chars=10**9)
    if "text_preview" in inspected:
        inspected["text"] = inspected.pop("text_preview")
    return inspected


def extract(path: Path, max_chars: int) -> dict[str, Any]:
    if not path.exists():
        return {"ok": False, "error": "file_not_found", "path": str(path)}
    suffix = path.suffix.lower()
    if suffix == ".docx":
        result = extract_docx(path)
    elif suffix == ".pptx":
        result = extract_pptx(path)
    elif suffix == ".xlsx":
        result = extract_xlsx(path)
    elif suffix == ".pdf":
        result = extract_pdf(path)
    else:
        result = {"ok": False, "error": "unsupported_extension", "extension": suffix}
    result["path"] = str(path)
    text = json.dumps(result, ensure_ascii=False, default=str)
    if len(text) > max_chars:
        result["truncated"] = True
        result["content_preview"] = text[:max_chars]
        for key in ("paragraphs", "tables", "slides", "sheets", "text"):
            result.pop(key, None)
    return result


def to_markdown(result: dict[str, Any]) -> str:
    lines = [f"# Document Extraction: {result.get('path', '')}", ""]
    lines.append("```json")
    lines.append(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    lines.append("```")
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract document content.")
    parser.add_argument("path")
    parser.add_argument("--format", choices=["json", "markdown"], default="json")
    parser.add_argument("--max-chars", type=int, default=50000)
    parser.add_argument("--no-log", action="store_true")
    args = parser.parse_args()

    result = extract(Path(args.path), args.max_chars)
    if args.format == "json":
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    else:
        print(to_markdown(result))
    if append_action and not args.no_log:
        append_action("agent", "doc-extract", f"{args.path}: ok={result.get('ok')}")
    return 0 if result.get("ok") else 2


if __name__ == "__main__":
    raise SystemExit(main())
