#!/usr/bin/env python3
"""Inspect Office and PDF files and emit concise JSON or Markdown summaries."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from log_action import append_action
except Exception:  # pragma: no cover - script should still run standalone
    append_action = None


def missing(module: str, package: str) -> dict[str, Any]:
    return {
        "ok": False,
        "missing_dependency": module,
        "install": f"Install {package} from SKILL/requirements.txt.",
    }


def inspect_docx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        import docx
    except ImportError:
        return missing("docx", "python-docx")
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append({"rows": len(rows), "columns": max((len(r) for r in rows), default=0), "preview": rows[:3]})
    text = "\n".join(paragraphs)
    return {
        "ok": True,
        "type": "docx",
        "paragraphs": len(paragraphs),
        "tables": len(tables),
        "text_preview": text[:max_chars],
        "table_preview": tables[:5],
    }


def inspect_pptx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError:
        return missing("pptx", "python-pptx")
    presentation = Presentation(str(path))
    slides = []
    collected: list[str] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        collected.extend(texts)
        slides.append({"slide": idx, "text_blocks": len(texts), "preview": texts[:5]})
    return {
        "ok": True,
        "type": "pptx",
        "slides": len(slides),
        "slide_preview": slides[:10],
        "text_preview": "\n".join(collected)[:max_chars],
    }


def inspect_xlsx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError:
        return missing("openpyxl", "openpyxl")
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
    sheets = []
    snippets: list[str] = []
    for ws in workbook.worksheets:
        rows = []
        for row in ws.iter_rows(max_row=5, values_only=True):
            values = ["" if value is None else str(value) for value in row]
            rows.append(values)
            snippets.append("\t".join(values))
        sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "preview": rows})
    workbook.close()
    return {
        "ok": True,
        "type": "xlsx",
        "sheets": sheets,
        "text_preview": "\n".join(snippets)[:max_chars],
    }


def inspect_pdf(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            import pdfplumber
        except ImportError:
            return missing("pypdf/pdfplumber", "pypdf or pdfplumber")
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
            return {"ok": True, "type": "pdf", "engine": "pdfplumber", "pages": len(pdf.pages), "text_preview": text[:max_chars]}
    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    metadata = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
    return {
        "ok": True,
        "type": "pdf",
        "engine": "pypdf",
        "pages": len(reader.pages),
        "metadata": metadata,
        "text_preview": text[:max_chars],
    }


def inspect(path: Path, max_chars: int) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if not path.exists():
        return {"ok": False, "error": "file_not_found", "path": str(path)}
    if suffix == ".docx":
        result = inspect_docx(path, max_chars)
    elif suffix == ".pptx":
        result = inspect_pptx(path, max_chars)
    elif suffix == ".xlsx":
        result = inspect_xlsx(path, max_chars)
    elif suffix == ".pdf":
        result = inspect_pdf(path, max_chars)
    else:
        result = {"ok": False, "error": "unsupported_extension", "extension": suffix}
    result["path"] = str(path)
    return result


def to_markdown(result: dict[str, Any]) -> str:
    lines = [f"# Document Inspection: {result.get('path', '')}", ""]
    for key, value in result.items():
        if key == "text_preview":
            continue
        lines.append(f"- {key}: {value}")
    if "text_preview" in result:
        lines.extend(["", "## Text Preview", "", str(result["text_preview"])])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect docx, pptx, xlsx, or pdf files.")
    parser.add_argument("path")
    parser.add_argument("--format", choices=["json", "markdown"], default="json")
    parser.add_argument("--max-chars", type=int, default=8000)
    parser.add_argument("--no-log", action="store_true")
    args = parser.parse_args()

    result = inspect(Path(args.path), args.max_chars)
    if args.format == "json":
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(to_markdown(result))
    if append_action and not args.no_log:
        append_action("agent", "doc-inspect", f"{args.path}: ok={result.get('ok')}")
    return 0 if result.get("ok") else 2


if __name__ == "__main__":
    raise SystemExit(main())
